"""Tests for reading alumni lists in different layouts (CH-100). All people are made up."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest
from sqlalchemy import select
from sqlalchemy.orm import Session

from careerhub.cli import main
from careerhub.core.migrate import migrate
from careerhub.core.models import Person
from careerhub.core.secrets import InMemoryKeyStore, new_database_key
from careerhub.core.storage import open_engine, session_factory
from careerhub.modules.network import import_alumni, preview_alumni
from careerhub.modules.network.layouts import SYNONYMS, normalize, read_table
from careerhub.profile import ProfileFileError


@pytest.fixture
def session(tmp_path: Path) -> Session:
    engine = open_engine(tmp_path / "t.db", new_database_key())
    migrate(engine)
    return session_factory(engine)()


def excel(path: Path, rows: list[list[object]], title: str = "Sheet1") -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = title
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    return path


def test_a_career_office_excel_file_with_a_title_row_and_odd_headers(
    tmp_path: Path, session: Session
) -> None:
    path = excel(
        tmp_path / "career.xlsx",
        [
            ["Example University Career Office: Physics Alumni, 2026"],
            [],
            [
                "Alumni Name",
                "Current Employer",
                "Job Title",
                "E-mail Address",
                "Graduation Year",
                "Notes",
            ],
            ["Ada Sample", "Acme Corp", "Engineer", "ada@example.com", 2019, "x"],
            ["Ben Example", "Globex", "Analyst", None, 2021, ""],
        ],
    )
    report = import_alumni(session, path, group="Example Physics")
    assert (report.created, report.skipped) == (2, 0)
    ada = session.scalars(select(Person).where(Person.full_name == "Ada Sample")).one()
    assert ada.how_known == "alumni: Example Physics" and ada.title == "Engineer"
    assert ada.email == "ada@example.com" and ada.notes == "class year: 2019"


def test_a_semicolon_csv_with_a_bom_and_no_group_column(tmp_path: Path, session: Session) -> None:
    path = tmp_path / "club.csv"
    path.write_bytes("﻿Name;Organisation;Position\nAda Sample;Acme;Eng\n".encode())
    assert import_alumni(session, path, group="Robotics Club").created == 1
    assert session.scalars(select(Person)).one().company.name == "Acme"  # type: ignore[union-attr]


def test_a_tab_separated_file_and_first_last_names(tmp_path: Path, session: Session) -> None:
    path = tmp_path / "list.tsv"
    path.write_text(
        "Given Name\tSurname\tEmployer\tClub\nEve\tSample\tAcme\tChess\n", encoding="utf-8"
    )
    assert import_alumni(session, path).created == 1
    assert session.scalars(select(Person)).one().full_name == "Eve Sample"


def test_a_windows_1252_file_is_read(tmp_path: Path, session: Session) -> None:
    path = tmp_path / "w.csv"
    path.write_bytes("Name,Company,Group\nJosé Garcia,Acmé,Club\n".encode("cp1252"))
    import_alumni(session, path)
    assert session.scalars(select(Person)).one().full_name == "José Garcia"


def test_the_first_sheet_is_used_unless_another_is_named(tmp_path: Path, session: Session) -> None:
    path = tmp_path / "two.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(["Name", "Company", "Group"])
    workbook.active.append(["First Sheet", "Acme", "G"])
    other = workbook.create_sheet("Second")
    other.append(["Name", "Company", "Group"])
    other.append(["Second Sheet", "Acme", "G"])
    workbook.save(path)
    import_alumni(session, path)
    import_alumni(session, path, sheet="Second")
    assert {p.full_name for p in session.scalars(select(Person))} == {"First Sheet", "Second Sheet"}
    with pytest.raises(ProfileFileError, match="no sheet named"):
        import_alumni(session, path, sheet="Nope")


def test_a_per_column_override_picks_the_right_column(tmp_path: Path, session: Session) -> None:
    path = tmp_path / "o.csv"
    path.write_text("Name,Employer,Company Size,Group\nAda,Acme,500,Club\n", encoding="utf-8")
    import_alumni(session, path, overrides={"company": "Company Size"})
    assert session.scalars(select(Person)).one().company.name == "500"  # type: ignore[union-attr]
    with pytest.raises(ProfileFileError, match="no column called"):
        import_alumni(session, path, overrides={"company": "Nope"})
    with pytest.raises(ProfileFileError, match="Unknown field"):
        import_alumni(session, path, overrides={"salary": "Employer"})


def test_missing_columns_are_named_in_the_error(tmp_path: Path, session: Session) -> None:
    path = tmp_path / "bad.csv"
    path.write_text("Name,Title\nAda,Eng\n", encoding="utf-8")
    with pytest.raises(ProfileFileError) as raised:
        import_alumni(session, path, group="G")
    assert "Company" in str(raised.value) and "Ada" not in str(raised.value)
    path.write_text("Name,Company\nAda,Acme\n", encoding="utf-8")
    with pytest.raises(ProfileFileError, match="Group"):
        import_alumni(session, path)


def test_rows_missing_name_or_company_are_skipped_and_the_summary_says_which(
    tmp_path: Path, session: Session
) -> None:
    path = tmp_path / "s.csv"
    path.write_text("Name,Company\n,Acme\nCy,\nDee,Acme\n", encoding="utf-8")
    report = import_alumni(session, path, group="G")
    assert (report.created, report.skipped, report.skipped_for) == (1, 2, ("company", "name"))


def test_the_same_person_twice_in_one_file_is_added_once(tmp_path: Path, session: Session) -> None:
    path = tmp_path / "d.csv"
    path.write_text("Name,Company\nAda Sample,Acme\nada sample,ACME Inc\n", encoding="utf-8")
    report = import_alumni(session, path, group="G")
    assert (report.created, report.unchanged) == (1, 1)


@pytest.mark.parametrize("text", ["", "\n\n", "just,some,words\n1,2,3\n", "x" * 10_000])
def test_unreadable_or_headerless_files_are_clear_errors(tmp_path: Path, text: str) -> None:
    path = tmp_path / "e.csv"
    path.write_text(text, encoding="utf-8")
    with pytest.raises(ProfileFileError) as raised:
        read_table(path)
    assert "header" in str(raised.value) and text[:20] not in str(raised.value) or text == ""


def test_a_damaged_excel_file_does_not_crash_or_echo(tmp_path: Path) -> None:
    path = tmp_path / "bad.xlsx"
    path.write_bytes(b"not really a zip file")
    with pytest.raises(ProfileFileError) as raised:
        read_table(path)
    assert "could not be read" in str(raised.value) and "not really" not in str(raised.value)


def test_the_synonym_table_is_normalized_and_has_no_overlaps() -> None:
    seen: dict[str, str] = {}
    for field_name, names in SYNONYMS.items():
        for name in names:
            assert name == normalize(name) and name not in seen, (field_name, name)
            seen[name] = field_name
    assert normalize("E-mail Address ") == "emailaddress"


def test_preview_describes_the_mapping_and_counts_without_saving_or_naming_anyone(
    tmp_path: Path, session: Session
) -> None:
    path = excel(
        tmp_path / "p.xlsx",
        [
            ["Alumni of Example"],
            ["Full Name", "Employer", "Hobby"],
            ["Zed Secret", "Acme", "chess"],
        ],
    )
    shown = preview_alumni(session, path, group="G")
    assert shown.header_row == 2 and shown.mapping == {"name": "Full Name", "company": "Employer"}
    assert shown.ignored == ("Hobby",) and shown.report.created == 1
    assert session.scalars(select(Person)).all() == []  # nothing saved


def test_the_command_previews_then_imports(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    monkeypatch.setenv("CAREERHUB_DATA_DIR", str(tmp_path))
    store = InMemoryKeyStore()
    path = tmp_path / "club.csv"
    path.write_text("Full Name,Employer,Hobby\nZed Secret,Acme,chess\n", encoding="utf-8")
    assert main(["alumni", "import", str(path), "--group", "Club", "--preview"], store) == 0
    out = capsys.readouterr().out
    assert "'Full Name' -> name" in out and "'Hobby' -> ignored" in out
    assert "Would add 1" in out and "Zed" not in out and "Acme" not in out
    assert main(["alumni", "import", str(path)], store) == 1  # no group yet
    assert "Group" in capsys.readouterr().out
    assert main(["people", "import-alumni", str(path), "--group", "Club"], store) == 0
    assert "1 added" in capsys.readouterr().out
    assert main(["alumni", "import", str(path), "--group", "Club", "--map", "oops"], store) == 2
    assert main(["alumni", "import", str(path), "--bogus"], store) == 2
