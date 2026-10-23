"""Tests for the document reader (CH-90). Sample files are generated here; none are real."""

from __future__ import annotations

import subprocess
from pathlib import Path

import docx
import pytest

from careerhub.profile import documents
from careerhub.profile.documents import MAX_FILE_BYTES, latex_to_text, read_document


def make_pdf(path: Path, pages: list[str]) -> Path:
    """Write a small PDF whose pages carry the given text lines ('' means a page with no text)."""
    objects: list[bytes] = []
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(len(pages)))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    font_id = 3 + 2 * len(pages)
    for i, text in enumerate(pages):
        stream = f"BT /F1 12 Tf 72 700 Td ({text}) Tj ET".encode() if text else b""
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {4 + 2 * i} 0 R "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode()
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref,
    )
    path.write_bytes(out)
    return path


def test_plain_text_and_markdown_are_read_as_they_are(tmp_path: Path) -> None:
    (tmp_path / "a.txt").write_text("Built a tool.\n\nLed a team.", encoding="utf-8")
    (tmp_path / "b.md").write_text("# Jane\n- Python", encoding="utf-8")
    assert read_document(tmp_path / "a.txt").text == "Built a tool.\n\nLed a team."
    assert read_document(tmp_path / "b.md").format == "md"


def test_latex_is_reduced_to_text(tmp_path: Path) -> None:
    (tmp_path / "r.tex").write_text(r"\textbf{Jane} built a tool. % private note", encoding="utf-8")
    result = read_document(tmp_path / "r.tex")
    assert result.text == "Jane built a tool." and not result.problems


def test_docx_paragraphs_and_tables_are_read(tmp_path: Path) -> None:
    document = docx.Document()
    document.add_paragraph("Software Engineer Intern, Acme")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "CS 101"
    table.rows[0].cells[1].text = "A"
    document.save(str(tmp_path / "r.docx"))
    result = read_document(tmp_path / "r.docx")
    assert "Software Engineer Intern, Acme" in result.text and "CS 101 | A" in result.text


def test_pdf_text_and_page_count(tmp_path: Path) -> None:
    path = make_pdf(
        tmp_path / "t.pdf", ["Course CS101 grade A in the fall term", "Second page text here ok"]
    )
    result = read_document(path)
    assert result.pages == 2 and not result.problems
    assert "CS101" in result.text and "Second page" in result.text
    assert result.chars == len(result.text)


def test_a_scanned_pdf_is_reported_not_returned_empty_silently(tmp_path: Path) -> None:
    result = read_document(make_pdf(tmp_path / "scan.pdf", ["", ""]))
    assert result.text == "" and result.pages == 2
    assert "scanned" in result.problems[0]


def test_pages_files_say_to_export_as_pdf(tmp_path: Path) -> None:
    (tmp_path / "r.pages").write_bytes(b"x")
    assert "export it as PDF" in read_document(tmp_path / "r.pages").problems[0]


@pytest.mark.parametrize("name", ["r.doc", "r.rtf", "r.odt"])
def test_mac_only_formats_are_reported_on_other_systems(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, name: str
) -> None:
    monkeypatch.setattr(documents.sys, "platform", "linux")
    (tmp_path / name).write_bytes(b"x")
    assert "only be read on a Mac" in read_document(tmp_path / name).problems[0]


def test_textutil_is_used_on_a_mac(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(documents.sys, "platform", "darwin")
    calls: list[list[str]] = []

    def fake_run(command: list[str], **kwargs: object) -> subprocess.CompletedProcess[bytes]:
        calls.append(command)
        return subprocess.CompletedProcess(command, 0, b"Line one\n\n  Line   two\n", b"")

    monkeypatch.setattr(documents.subprocess, "run", fake_run)
    (tmp_path / "r.rtf").write_bytes(b"x")
    assert read_document(tmp_path / "r.rtf").text == "Line one\nLine two"
    assert calls[0][:3] == ["textutil", "-convert", "txt"]


def test_a_failed_textutil_run_is_a_problem(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(documents.sys, "platform", "darwin")
    monkeypatch.setattr(
        documents.subprocess, "run", lambda c, **k: subprocess.CompletedProcess(c, 1, b"", b"")
    )
    (tmp_path / "r.doc").write_bytes(b"x")
    assert "exit 1" in read_document(tmp_path / "r.doc").problems[0]


def test_bad_inputs_are_problems_that_never_quote_the_text(tmp_path: Path) -> None:
    secret = "SECRET-RESUME-TEXT"
    (tmp_path / "broken.pdf").write_text(secret)
    (tmp_path / "broken.docx").write_text(secret)
    (tmp_path / "r.xyz").write_text(secret)
    for name in ("broken.pdf", "broken.docx", "r.xyz", "missing.txt"):
        result = read_document(tmp_path / name)
        assert result.text == "" and result.problems
        assert secret not in " ".join(result.problems)


def test_a_huge_file_is_refused(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path = tmp_path / "big.txt"
    path.write_text("x")
    monkeypatch.setattr(documents, "MAX_FILE_BYTES", 0)
    assert MAX_FILE_BYTES > 0 and "larger than" in read_document(path).problems[0]


def test_invalid_utf8_and_null_bytes_do_not_crash(tmp_path: Path) -> None:
    (tmp_path / "a.txt").write_bytes(b"ok \xff\xfe\x00 fine")
    assert read_document(tmp_path / "a.txt").text.startswith("ok")
    assert "\x00" not in read_document(tmp_path / "a.txt").text


def test_empty_files_give_empty_text(tmp_path: Path) -> None:
    (tmp_path / "a.txt").write_text("")
    assert read_document(tmp_path / "a.txt").chars == 0


@pytest.mark.parametrize(
    ("source", "text"),
    [
        (r"Jos\'e Garc\'ia", "José García"),
        (r"Zo\"e M\"uller", "Zoë Müller"),
        (r"Fran\c{c}ois", "François"),
        (r"\v{S}koda and \'{E}milie", "Škoda and Émilie"),
        (r"Stra\ss e and \o{}ystein", "Straße and øystein"),
    ],
)
def test_latex_accents_become_real_letters_so_names_can_be_matched(source: str, text: str) -> None:
    from careerhub.profile.redact import personal_text

    assert latex_to_text(source) == text
    assert "[name]" in personal_text(latex_to_text(source), text)


def test_a_partly_scanned_pdf_keeps_the_text_pages_and_names_the_scanned_ones(
    tmp_path: Path,
) -> None:
    """CH-118: decided per page, not by the document's average."""
    path = make_pdf(tmp_path / "t.pdf", ["Course CS101 grade A in the fall term", "", ""])
    result = read_document(path)
    assert "CS101" in result.text and result.pages == 3
    assert result.problems == ("Pages 2 and 3 look scanned and were skipped (no text layer).",)


def test_one_scanned_page_among_text_pages_is_named_in_the_singular(tmp_path: Path) -> None:
    path = make_pdf(
        tmp_path / "t.pdf", ["First page has plenty of text", "", "Third page has text too ok"]
    )
    result = read_document(path)
    assert "First page" in result.text and "Third page" in result.text
    assert result.problems == ("Page 2 looks scanned and was skipped (no text layer).",)


_TEXT_BOX = (
    '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
    "<mc:AlternateContent><mc:Choice><w:txbxContent><w:p><w:r><w:t>Skills: Python, SQL</w:t></w:r>"
    "</w:p></w:txbxContent></mc:Choice><mc:Fallback><w:txbxContent><w:p><w:r>"
    "<w:t>Skills: Python, SQL</w:t></w:r></w:p></w:txbxContent></mc:Fallback>"
    "</mc:AlternateContent></w:r>"
)


def test_docx_text_boxes_and_nested_tables_are_read_once(tmp_path: Path) -> None:
    """CH-118: some resume templates keep skills in a text box or a table inside a table."""
    from docx.oxml import parse_xml

    document = docx.Document()
    paragraph = document.add_paragraph("Jane Sample, Software Engineer")
    paragraph._p.append(parse_xml(_TEXT_BOX))
    outer = document.add_table(rows=1, cols=1)
    cell = outer.rows[0].cells[0]
    cell.text = "Experience"
    inner = cell.add_table(rows=1, cols=2)
    inner.rows[0].cells[0].text = "Acme Corp"
    inner.rows[0].cells[1].text = "2024"
    document.sections[0].header.paragraphs[0].text = "jane@example.com 555 123 4567"
    document.save(str(tmp_path / "r.docx"))
    text = read_document(tmp_path / "r.docx").text
    assert text.count("Skills: Python, SQL") == 1  # the fallback copy is skipped
    assert "Acme Corp | 2024" in text and "Experience" in text
    assert "jane@example.com" not in text  # headers stay skipped
