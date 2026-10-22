"""Read an alumni or contact list in whatever layout it comes in (CH-100).

The user does not know in advance what an export looks like, so nothing about the file is
hard-coded: the delimiter is sniffed, `.xlsx` files are read from their first sheet (or `--sheet`),
title rows above the header are skipped, and columns are matched by *synonyms* ("Employer",
"Current Company", "Organisation" are all the company), ignoring case, spacing, and punctuation. The
synonym table is one data structure, `SYNONYMS`, easy to extend. Only name and company are required.
Nothing here prints or logs a person's details.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from pathlib import Path

from careerhub.profile.files import ProfileFileError

#: Field -> the column names that mean it (already normalized: lowercase letters and digits only).
SYNONYMS: dict[str, tuple[str, ...]] = {
    "name": ("name", "fullname", "alumniname", "contact", "contactname", "person"),
    "first": ("first", "firstname", "givenname", "forename"),
    "last": ("last", "lastname", "surname", "familyname"),
    "company": (
        "company", "employer", "currentemployer", "organization", "organisation", "firm",
        "currentcompany", "workplace", "companyname", "employername",
    ),
    "title": ("title", "jobtitle", "position", "role", "currenttitle", "currentposition"),
    "email": ("email", "emailaddress", "mail", "primaryemail"),
    "url": ("url", "linkedin", "linkedinurl", "profile", "profileurl", "linkedinprofile"),
    "group": ("group", "club", "program", "programme", "organizationtype", "affiliation"),
    "school": ("school", "college", "university", "almamater", "institution", "alumniof"),
    "class_year": ("classyear", "graduationyear", "gradyear", "class", "year", "cohort"),
    "major": ("major", "degree", "concentration", "fieldofstudy"),
    "location": ("location", "city", "region", "metro"),
}  # fmt: skip
EXTRA_FIELDS = ("class_year", "major", "location")  # kept in a new person's notes
HEADER_SEARCH_ROWS = 30
MAX_ROWS = 50_000


def normalize(header: str) -> str:
    """Lowercase and keep letters and digits only, so "E-mail Address" matches "emailaddress"."""
    return re.sub(r"[^a-z0-9]", "", header.lower())


@dataclass
class Table:
    """The cells of a list, from its header row down."""

    header: list[str]
    rows: list[list[str]]
    header_row: int  # 1-based position in the file, for the preview
    mapping: dict[str, int] = field(default_factory=dict)  # field -> column index
    ignored: list[str] = field(default_factory=list)  # headers that matched nothing

    def get(self, row: list[str], name: str) -> str:
        """Return the cell for field `name` in `row` (empty if unmapped or short)."""
        index = self.mapping.get(name)
        return row[index].strip() if index is not None and index < len(row) else ""

    @property
    def has_name(self) -> bool:
        """Say whether a person's name can be built (a name column, or first plus last)."""
        return "name" in self.mapping or {"first", "last"} <= self.mapping.keys()

    def person_name(self, row: list[str]) -> str:
        """Return the person's full name from the name column or first and last."""
        full = self.get(row, "name")
        return full or " ".join(f"{self.get(row, 'first')} {self.get(row, 'last')}".split())


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _read_delimited(path: Path) -> list[list[str]]:
    text = _decode(path.read_bytes())
    sample = text[:8192]
    if path.suffix.lower() == ".tsv":
        delimiter = "\t"
    else:
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    rows: list[list[str]] = []
    for row in reader:
        rows.append(row)
        if len(rows) > MAX_ROWS:
            break
    return rows


def _read_excel(path: Path, sheet: str | None) -> list[list[str]]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet is not None and sheet not in workbook.sheetnames:
            raise ProfileFileError(f"{path.name} has no sheet named {sheet!r}.")
        worksheet = workbook[sheet] if sheet else workbook.worksheets[0]
        rows: list[list[str]] = []
        for values in worksheet.iter_rows(values_only=True):
            rows.append(["" if cell is None else str(cell) for cell in values])
            if len(rows) > MAX_ROWS:
                break
        return rows
    finally:
        workbook.close()


def _score(row: list[str]) -> int:
    """Count the cells of `row` that look like known column names."""
    known = {name for names in SYNONYMS.values() for name in names}
    return sum(normalize(cell) in known for cell in row)


def read_table(
    path: Path, sheet: str | None = None, overrides: dict[str, str] | None = None
) -> Table:
    """Read `path` into a `Table` with its columns mapped. Raises `ProfileFileError` if unreadable.

    `overrides` maps a field to the exact header to use for it (for example
    `{"company": "Current Employer"}`).
    """
    if not path.exists():
        raise ProfileFileError(f"{path.name} does not exist.")
    try:
        if path.suffix.lower() in (".xlsx", ".xlsm"):
            cells = _read_excel(path, sheet)
        else:
            cells = _read_delimited(path)
    except ProfileFileError:
        raise
    except Exception as error:  # a damaged file must not crash; the error carries no content
        raise ProfileFileError(f"{path.name} could not be read ({type(error).__name__}).") from None
    top = cells[:HEADER_SEARCH_ROWS]
    scored = [(_score(row), index) for index, row in enumerate(top)]
    best = next(((s, i) for s, i in scored if s >= 2), None) or next(
        ((s, i) for s, i in scored if s >= 1), None
    )
    if best is None:
        raise ProfileFileError(
            f"{path.name}: no header row was found (looked at the first 30 rows)."
        )
    index = best[1]
    header = [cell.strip() for cell in cells[index]]
    rows = [row for row in cells[index + 1 :] if any(cell.strip() for cell in row)]
    table = Table(header, rows, header_row=index + 1)
    _map_columns(table, overrides or {})
    return table


def _map_columns(table: Table, overrides: dict[str, str]) -> None:
    """Match each header to a field by synonym (first match wins), honoring `overrides`."""
    by_header = {normalize(h): i for i, h in reversed(list(enumerate(table.header))) if h}
    for field_name, header in overrides.items():
        if field_name not in SYNONYMS:
            raise ProfileFileError(
                f"Unknown field {field_name!r}; use one of {', '.join(SYNONYMS)}."
            )
        if normalize(header) not in by_header:
            raise ProfileFileError(f"There is no column called {header!r} in the file.")
        table.mapping[field_name] = by_header[normalize(header)]
    taken = set(table.mapping.values())
    for index, header in enumerate(table.header):
        if index in taken or not header:
            continue
        for field_name, names in SYNONYMS.items():
            if field_name not in table.mapping and normalize(header) in names:
                table.mapping[field_name] = index
                break
        else:
            table.ignored.append(header)
