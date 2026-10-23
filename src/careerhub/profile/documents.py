"""Read the text of a resume or transcript in any common format (CH-90).

`read_document(path)` returns the plain text plus what is known about the file (format, page
count, character count) and a list of **problems**. It never raises for a bad file and never puts
document text into a problem, a log, or an error, because the text is Personal data.

Formats: `.tex`, `.pdf` (pypdfium2), `.docx` (python-docx), `.txt`, `.md`; `.doc`, `.rtf`, `.odt`
through macOS's built-in `textutil` (on other systems they are reported as Mac-only); `.pages` is
reported as "export it as PDF". A PDF with almost no text layer is reported as scanned, because it
needs text recognition (a follow-up that runs on the Mac).
"""

from __future__ import annotations

import re
import subprocess
import sys
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Any

#: Files bigger than this are refused, so a huge file can't use all memory.
MAX_FILE_BYTES = 20 * 1024 * 1024
#: Fewer characters per page than this means the PDF has no real text layer (a scan).
SCANNED_CHARS_PER_PAGE = 20
TEXTUTIL_FORMATS = {".doc", ".rtf", ".odt"}
SUPPORTED = (".tex", ".pdf", ".docx", ".txt", ".md", *sorted(TEXTUTIL_FORMATS))


_COMMENT = re.compile(r"(?<!\\)%.*")
_ENVIRONMENT = re.compile(r"\\(?:begin|end)\{[^}]*\}")
_COMMAND_WITH_ARG = re.compile(r"\\[a-zA-Z]+\*?(?:\[[^\]]*\])?\{([^{}]*)\}")
_BARE_COMMAND = re.compile(r"\\[a-zA-Z]+\*?")
_ESCAPED = re.compile(r"\\([_&%#$])")


_ACCENTS = {
    "`": "\u0300", "'": "\u0301", "^": "\u0302", "~": "\u0303", "=": "\u0304", "u": "\u0306",
    ".": "\u0307", '"': "\u0308", "r": "\u030a", "H": "\u030b", "v": "\u030c", "c": "\u0327",
    "k": "\u0328",
}  # fmt: skip
_ACCENT = re.compile(r"\\(?:([`'^~=.\"])\s*\{?\\?([A-Za-z])\}?|([uvHcrk])\s*\{([A-Za-z])\})")
_SPECIAL = {r"\ss": "ß", r"\o": "ø", r"\O": "Ø", r"\ae": "æ", r"\AE": "Æ", r"\l": "ł"}


def _accented(match: re.Match[str]) -> str:
    """Turn a LaTeX accent command (acute e, cedilla c...) into the accented letter."""
    mark = match.group(1) or match.group(3)
    letter = match.group(2) or match.group(4)
    return unicodedata.normalize("NFC", letter + _ACCENTS[mark])


def latex_to_text(source: str) -> str:
    """Reduce LaTeX source to readable text: comments, commands, and braces removed."""
    text = _ESCAPED.sub(r"\1", _COMMENT.sub("", source))
    text = _ACCENT.sub(_accented, text)
    for command, letter in _SPECIAL.items():
        text = re.sub(re.escape(command) + r"(?![A-Za-z])\s?", letter, text)
    text = _ENVIRONMENT.sub("\n", text)
    for _ in range(5):  # unwrap nested \textbf{\emph{...}}
        text = _COMMAND_WITH_ARG.sub(r"\1", text)
    text = _BARE_COMMAND.sub(" ", text).replace("\\\\", "\n")
    text = re.sub(r"[{}]", "", text)
    lines = (" ".join(line.split()) for line in text.splitlines())
    return "\n".join(line for line in lines if line)


@dataclass(frozen=True)
class DocumentText:
    """What was read from one file. `text` is empty when `problems` says why."""

    text: str
    format: str
    pages: int | None
    problems: tuple[str, ...] = ()

    @property
    def chars(self) -> int:
        """The number of characters of text (safe to show; the text itself is not)."""
        return len(self.text)


def read_document(path: Path) -> DocumentText:
    """Return the text of `path`, or an empty text with the reasons in `problems`."""
    suffix = path.suffix.lower()
    kind = suffix.lstrip(".") or "unknown"

    def fail(message: str) -> DocumentText:
        return DocumentText("", kind, None, (message,))

    if suffix == ".pages":
        return fail("Pages files can't be read: export it as PDF from Pages and use that.")
    if suffix not in SUPPORTED:
        return fail(
            f"Unsupported format {suffix or '(none)'!r}. Use one of: {', '.join(SUPPORTED)}."
        )
    try:
        if not path.is_file():
            return fail("The file does not exist.")
        if path.stat().st_size > MAX_FILE_BYTES:
            return fail("The file is larger than 20 MB.")
        if suffix == ".pdf":
            return _read_pdf(path)
        if suffix == ".docx":
            return _read_docx(path)
        if suffix in TEXTUTIL_FORMATS:
            return _read_with_textutil(path, kind)
        raw = path.read_text(encoding="utf-8", errors="replace")
        if suffix == ".tex":
            raw = latex_to_text(raw)
        return DocumentText(raw.replace("\x00", ""), kind, None)
    except Exception as error:  # a damaged file must not crash; only the error's type is kept
        return fail(f"The file could not be read ({type(error).__name__}).")


def _page_list(numbers: list[int]) -> str:
    """Write page numbers for a message: "page 2", "pages 2 and 3", "pages 1, 2 and 4"."""
    if len(numbers) == 1:
        return f"page {numbers[0]}"
    return f"pages {', '.join(map(str, numbers[:-1]))} and {numbers[-1]}"


def _scanned_message(numbers: list[int]) -> str:
    """Say which pages were skipped as scanned."""
    verb = "looks scanned and was" if len(numbers) == 1 else "look scanned and were"
    return f"{_page_list(numbers).capitalize()} {verb} skipped (no text layer)."


def _read_pdf(path: Path) -> DocumentText:
    """Extract the text layer of a PDF, page by page.

    A page with almost no text is scanned (an image). If every page is, nothing is returned; if only
    some are, the text pages are returned and the problem names the skipped ones (CH-118).
    """
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(path))
    try:
        pages = len(pdf)
        parts: list[str] = []
        for index in range(pages):
            page = pdf[index]
            textpage = page.get_textpage()
            parts.append(textpage.get_text_range())
            textpage.close()
            page.close()
    finally:
        pdf.close()
    cleaned = ["\n".join(" ".join(line.split()) for line in part.splitlines()) for part in parts]
    scanned = [
        number
        for number, text in enumerate(cleaned, start=1)
        if len(text.replace("\n", "")) < SCANNED_CHARS_PER_PAGE
    ]
    if pages and len(scanned) == pages:
        return DocumentText(
            "",
            "pdf",
            pages,
            ("This PDF looks scanned (no text layer): it needs text recognition.",),
        )
    text = "\n".join(
        t for number, t in enumerate(cleaned, start=1) if number not in scanned
    ).strip()
    problems = (_scanned_message(scanned),) if scanned else ()
    return DocumentText(text, "pdf", pages, problems)


_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"


def _table_lines(table: Any) -> list[str]:
    """Return a table's rows as "cell | cell" lines, then the rows of tables nested in its cells."""
    lines: list[str] = []
    nested: list[Any] = []
    for row in table.rows:
        lines.append(" | ".join(cell.text for cell in row.cells))
        for cell in row.cells:
            nested.extend(cell.tables)
    for inner in nested:
        lines.extend(_table_lines(inner))
    return lines


def _text_box_lines(document: Any) -> list[str]:
    """Return the paragraphs inside text boxes (some resume templates keep skills there).

    Word stores each text box twice (a modern copy and a `Fallback` copy); the fallback is skipped.
    Headers and footers are not read: they are usually contact details.
    """
    lines = []
    for box in document.element.body.iter(f"{{{_W}}}txbxContent"):
        if any(a.tag == f"{{{_MC}}}Fallback" for a in box.iterancestors()):
            continue
        for paragraph in box.iter(f"{{{_W}}}p"):
            lines.append("".join(t.text or "" for t in paragraph.iter(f"{{{_W}}}t")))
    return lines


def _read_docx(path: Path) -> DocumentText:
    """Extract paragraphs, table cells (nested tables too), and text boxes from a Word file."""
    import docx

    document = docx.Document(str(path))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        lines.extend(_table_lines(table))
    lines.extend(_text_box_lines(document))
    text = "\n".join(" ".join(line.split()) for line in lines if line.strip())
    return DocumentText(text, "docx", None)


def _read_with_textutil(path: Path, kind: str) -> DocumentText:
    """Convert .doc, .rtf, or .odt with macOS's `textutil`."""
    if sys.platform != "darwin":
        return DocumentText(
            "",
            kind,
            None,
            (f".{kind} files can only be read on a Mac (save it as .docx or PDF instead).",),
        )
    done = subprocess.run(
        ["textutil", "-convert", "txt", "-stdout", str(path)],
        capture_output=True,
        timeout=60,
        check=False,
    )
    if done.returncode != 0:
        return DocumentText(
            "", kind, None, (f"textutil could not read the file (exit {done.returncode}).",)
        )
    text = done.stdout.decode("utf-8", errors="replace").replace("\x00", "")
    return DocumentText(
        "\n".join(" ".join(row.split()) for row in text.splitlines() if row.strip()), kind, None
    )
